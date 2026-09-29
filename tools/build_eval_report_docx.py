from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "docs" / "reports"
DATA_PATH = REPORT_DIR / "evaluation-data.json"
OUT_PATH = REPORT_DIR / "plum-evaluation-report-google-docs.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None, size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = tr_pr.find(qn("w:tblHeader"))
    if repeat is None:
        repeat = OxmlElement("w:tblHeader")
        tr_pr.append(repeat)
    repeat.set(qn("w:val"), "true")


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_meta_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "Field", bold=True, color="FFFFFF")
    set_cell_text(hdr[1], "Value", bold=True, color="FFFFFF")
    set_cell_shading(hdr[0], "404040")
    set_cell_shading(hdr[1], "404040")
    set_cell_width(hdr[0], 1.7)
    set_cell_width(hdr[1], 4.8)
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(cells[1], value)
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_cell_width(cells[0], 1.7)
        set_cell_width(cells[1], 4.8)


def add_case_table(doc: Document, records: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ["Case", "Scenario", "Expected", "Produced", "Amount", "Match"]
    widths = [0.62, 2.35, 1.0, 1.0, 0.88, 0.65]
    for i, label in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], label, bold=True, color="FFFFFF", size=8.5)
        set_cell_shading(table.rows[0].cells[i], "1F4E79")
        set_cell_width(table.rows[0].cells[i], widths[i])
    repeat_header_row(table.rows[0])
    keep_row_together(table.rows[0])
    for index, record in enumerate(records, start=1):
        cells = table.add_row().cells
        output = record["output"]
        expected = record["expected"].get("decision")
        produced = output.get("decision")
        decision_labels = {
            None: "None",
            "APPROVED": "Approved",
            "REJECTED": "Rejected",
            "PARTIAL": "Partial",
            "MANUAL_REVIEW": "Review",
        }
        values = [
            record["case_id"],
            record["case_name"].replace("Pre-Authorization", "Pre-Auth").replace("Multiple Same-Day", "Same-Day"),
            decision_labels.get(expected, str(expected)),
            decision_labels.get(produced, str(produced)),
            "None" if output.get("approved_amount") is None else f"Rs {output['approved_amount']:,}",
            "Yes" if record["matched"] else "No",
        ]
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=8.5)
            set_cell_width(cells[i], widths[i])
            if index % 2 == 0:
                set_cell_shading(cells[i], "F3F7FB")
        for i in (0, 2, 3, 4, 5):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        keep_row_together(table.rows[-1])


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_case_summaries(doc: Document, records: list[dict]) -> None:
    for record in records:
        output = record["output"]
        heading = doc.add_heading(f"{record['case_id']} {record['case_name']}", level=3)
        heading.paragraph_format.keep_with_next = True
        decision = output.get("decision") or output.get("state")
        amount = output.get("approved_amount")
        p = doc.add_paragraph()
        p.add_run("Result: ").bold = True
        p.add_run(
            f"{decision}; approved amount "
            f"{'not applicable' if amount is None else 'Rs ' + format(amount, ',')}; "
            f"match {'yes' if record['matched'] else 'no'}."
        )
        reasons = output.get("reasons") or output.get("correction_requests") or []
        if reasons:
            reason_text = "; ".join(
                str(item.get("code") or item.get("message") or item) for item in reasons[:4]
            )
            p = doc.add_paragraph()
            p.add_run("Primary evidence signal: ").bold = True
            p.add_run(reason_text)
        checks = record.get("behavior_checks", {})
        p = doc.add_paragraph()
        p.add_run("Explicit behavior checks: ").bold = True
        p.add_run("passed" if checks.get("matched") else "failed")


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 10.5)]:
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Title"
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def main() -> None:
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = data["records"]
    fingerprints = data["fingerprints"]
    passed = sum(1 for record in records if record["matched"])

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Plum Claims Evaluation Report")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.add_run("Google Docs ready formatted version of the policy pipeline evaluation").italic = True
    doc.add_paragraph(f"Prepared from generated evaluation artifacts on {date.today().isoformat()}.")

    doc.add_heading("Executive Result", level=1)
    p = doc.add_paragraph()
    p.add_run("Conclusion: ").bold = True
    p.add_run(
        f"{passed} of {len(records)} supplied structured fixture cases matched the expected decision, amount, reason, confidence rule, and explicit behavior checks."
    )
    p = doc.add_paragraph()
    p.add_run("Scope: ").bold = True
    p.add_run(
        "This evaluation proves deterministic policy pipeline behavior against structured fixtures. It does not prove OCR accuracy, handwriting support, multilingual extraction, image quality handling, or calibrated confidence."
    )

    add_meta_table(
        doc,
        [
            ("Policy", "PLUM_GHI_2024"),
            ("Cases", str(len(records))),
            ("Result", f"{passed}/{len(records)} matched"),
            ("Policy file sha256", fingerprints["policy_sha256"]),
            ("Canonical policy sha256", fingerprints["policy_canonical_sha256"]),
            ("Fixture file sha256", fingerprints["fixture_sha256"]),
            ("Full machine output", "docs/reports/evaluation-data.json"),
        ],
    )

    doc.add_heading("Case Results", level=1)
    add_case_table(doc, records)

    doc.add_heading("Interpretation and Limits", level=1)
    add_bullets(
        doc,
        [
            "The engine reads only the canonical policy produced from the unmodified policy file; every interpretation, merge, and conflict resolution is captured in the canonical audit trail and referenced from decision traces.",
            "The per-claim ceiling is the maximum of the global per-claim limit and the relevant category sub-limit, tested on eligible amount after excluded lines are removed. A matched pre-authorization rule governs amounts above the ceiling instead.",
            "Category sub-limits are annual per-member caps on net benefit for the category's own service lines. For consultation, tests and medicines billed with a consultation are not treated as consultation service lines.",
            "Aggregate limits are applied when utilisation accompanies the claim. If utilisation is absent, those rules are recorded as not evaluated, disclosed as advisory reasons on payable outcomes, and reflected in confidence.",
            "The supplied cases do not carry a submission date, so the 30-day submission deadline is not evaluated for them. Web intake stamps a server-side submission date at intake.",
            "Confidence values are heuristic evidence-completeness scores, not calibrated probabilities. TC011 intentionally simulates optional risk enrichment failure and records the degraded confidence and review recommendation.",
        ],
    )

    doc.add_heading("Case Output Summaries", level=1)
    add_case_summaries(doc, records)

    doc.add_heading("Generated Source Files", level=1)
    add_bullets(
        doc,
        [
            "docs/reports/evaluation.md contains the complete Markdown report with trace excerpts.",
            "docs/reports/evaluation-data.json contains the full structured output for every case.",
            "docs/reports/verification-summary.md contains the overall verification summary generated by make verify.",
        ],
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
